from pathlib import Path
import csv
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'report'/'Healthcare_Analytics_Report_Working_Draft.docx'
NAVY='14263D'; BLUE='3A86FF'; CORAL='EF476F'; GRAY='586574'; PALE='EEF4FA'; RED='C8102E'

def shade(cell, fill):
    pr=cell._tc.get_or_add_tcPr(); el=OxmlElement('w:shd'); el.set(qn('w:fill'),fill); pr.append(el)

def border(cell):
    pr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders'); pr.append(b)
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'5'); x.set(qn('w:color'),'D9E0E8'); b.append(x)

def celltext(cell, text, bold=False, color='26364A', size=8.5):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None, size=8.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for j,h in enumerate(headers):
        celltext(t.rows[0].cells[j],h,True,'FFFFFF',size); shade(t.rows[0].cells[j],NAVY); border(t.rows[0].cells[j])
    trpr=t.rows[0]._tr.get_or_add_trPr(); mark=OxmlElement('w:tblHeader'); mark.set(qn('w:val'),'true'); trpr.append(mark)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,v in enumerate(row):
            celltext(cells[j],v,False,'26364A',size)
            if i%2: shade(cells[j],'F5F8FB')
            border(cells[j])
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths): row.cells[j].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def prompt(doc, title, lines=2):
    # Clearly separated, editable student-response area without long underscore rules.
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'F1F6FC'); border(c); c.text=''
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(4); r=p.add_run('STUDENT RESPONSE  |  '+title)
    r.bold=True; r.font.name='Arial'; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(BLUE)
    q=c.add_paragraph('Write your interpretation here in your own words.'); q.paragraph_format.space_after=Pt(2)
    for _ in range(max(1,min(lines,3))):
        q=c.add_paragraph(); q.paragraph_format.space_after=Pt(3); q.paragraph_format.line_spacing=1.05
        q.add_run(' ').font.size=Pt(9)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def h(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; p.add_run(text); return p

def fact(doc,label,value):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); a=p.add_run(label+': '); a.bold=True; a.font.color.rgb=RGBColor.from_string(NAVY); p.add_run(value)

def narrative(doc,text):
    p=doc.add_paragraph(text); p.paragraph_format.keep_together=True; p.paragraph_format.space_after=Pt(7); return p

def image(doc,name,width=6.0,caption=''):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    p.add_run().add_picture(str(ROOT/'outputs'/'plots'/name),width=Inches(width))
    if caption:
        q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(3)
        r=q.add_run(caption); r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

def csvrows(name):
    with (ROOT/'outputs'/name).open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string('26364A'); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
for name,size in [('Title',23),('Heading 1',15),('Heading 2',11.5)]:
    s=doc.styles[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string('000000'); s.paragraph_format.space_before=Pt(12 if name=='Heading 1' else 7); s.paragraph_format.space_after=Pt(5)
head=sec.header.paragraphs[0]; head.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=head.add_run('SIT NAGPUR  |  HEALTHCARE ANALYTICS'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(GRAY)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=foot.add_run('OASIS cognitive status project  |  '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(GRAY)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)

# Cover in the formal college heading style used in the previous project.
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(30)
r=p.add_run('SYMBIOSIS INSTITUTE OF TECHNOLOGY, NAGPUR'); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=RGBColor.from_string(RED)
for text,size,color,bold in [('Symbiosis International (Deemed University)',13,'202020',True),('(Established under section 3 of the UGC Act, 1956)',8,GRAY,False),('Re-accredited by NAAC with A++ Grade  |  Awarded Category I by UGC',8,GRAY,False),('Founder: Prof. Dr. S. B. Mujumdar, M.Sc., Ph.D.',9,'1478B8',True)]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2); r=p.add_run(text); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); r.bold=bold
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); r=p.add_run('HEALTHCARE ANALYTICS MINI PROJECT'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Cognitive Status Classification Using OASIS-1 Clinical and MRI Measures')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Editable report working draft'); r.italic=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(GRAY)
for text in ['Submitted by','Atharv Jaiswal  |  PRN: 23070521023','Shreya Bansal  |  PRN: 23070521144','Healthcare Analytics mentors: Dr. Bhupesh and Dr. Prerna','October 2026']:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text); r.font.size=Pt(10 if text=='Submitted by' else 11); r.bold=text.startswith(('Atharv','Shreya'))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); r=p.add_run('AI-assisted explanatory draft. Verify Moodle dataset approval, check every result against the workbook, and follow course policy before submission.'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(CORAL)

h(doc,'Report guide',1)
doc.add_paragraph('This report presents a cross-sectional classification analysis of recorded Clinical Dementia Rating (CDR) groups in the supplied OASIS-1 workbook. It documents the dataset, preprocessing, model comparison and observed test results. The shaded response boxes are for your own interpretation; complete them after checking the source data and rerunning the analysis. The course requires original student writing.')
h(doc,'Abstract',1); narrative(doc,'This project evaluates whether demographic characteristics and MRI-derived measures can distinguish OASIS-1 records with a recorded Clinical Dementia Rating (CDR) of 0 from records with CDR above 0. The workbook contains 436 rows and 12 attributes; CDR is recorded for 235 rows, while 201 rows have no target label and are excluded from supervised modeling. The analysis compares logistic regression, a shallow decision tree, K-nearest neighbours, and a soft-voting ensemble across demographic, MRI-augmented, and MMSE-augmented feature sets. Thresholds are selected on validation data using F2 and evaluated on a separate 47-row test set. The ensemble reached a test ROC-AUC of 0.893, while the decision tree reached the highest test F2 of 0.893; KNN had the highest test accuracy at 0.766. These results are exploratory and split-dependent. Because the data are cross-sectional and MMSE is measured contemporaneously with CDR, the analysis does not estimate future dementia risk or support clinical decisions.')
h(doc,'Introduction',1); narrative(doc,'Healthcare analytics can describe patterns in clinical and imaging measurements and assess how well those measurements separate recorded patient groups. This project uses the CDR grouping documented in a cross-sectional OASIS-1 dataset. Its purpose is to compare analytical approaches and inspect their errors, not to diagnose individuals. A model trained to classify a contemporaneous label does not show that it can predict a later outcome.')
h(doc,'Problem statement',1)
fact(doc,'Outcome','Recorded CDR = 0 versus recorded CDR > 0; rows with missing CDR are excluded.')
fact(doc,'Primary question','Do demographic and MRI-derived measures distinguish the two recorded CDR groups, and how does adding MMSE change test performance?')
narrative(doc,'The practical question is whether the available demographic, MRI summary, and cognitive-score variables contain information that distinguishes the two recorded CDR groups in held-out records. Since the workbook provides cross-sectional observations rather than a future outcome or longitudinal prediction target, the analysis can answer only a classification question about recorded status.')

h(doc,'Objectives',1)
for item in ['Describe target availability, missingness, repeated sessions, and class balance.','Compare demographic, MRI-augmented, and MMSE-augmented feature sets.','Compare logistic regression, a shallow decision tree, K-nearest neighbours, and a soft-voting ensemble with a majority-class baseline.','Select thresholds on validation data using F2, then report held-out test metrics.','Identify limits on interpretation and generalization.']: doc.add_paragraph(item,style='List Bullet')
narrative(doc,'The objectives are to audit target completeness and class balance; compare performance as feature groups are added; evaluate multiple classifiers against simple baselines; select operating thresholds without using the test set; and describe the practical limits of the resulting estimates.')
h(doc,'Literature review and related work',1)
fact(doc,'Dataset paper','Marcus et al. (2007), Journal of Cognitive Neuroscience 19(9), 1498-1507. DOI 10.1162/jocn.2007.19.9.1498.')
narrative(doc,'The OASIS dataset paper by Marcus and colleagues introduces an open-access collection of cross-sectional MRI data spanning young, middle-aged, nondemented, and demented older adults. This project uses the supplied tabular summary rather than raw MRI images and focuses on recorded CDR grouping. Its question and analysis are narrower than the broader dataset description. Add and discuss other related studies only after reading them directly and confirming they are appropriate for the course.')
h(doc,'Dataset description',1)
table(doc,['Workbook item','Observed value'],[['Sheet','oasis_cross-sectional'],['Rows and attributes','436 rows, 12 columns'],['Participant IDs','416 unique participant prefixes across 436 MRI session rows'],['Known CDR','235 rows: 135 at CDR 0 and 100 above 0'],['Missing CDR','201 rows excluded from target modeling'],['Known-target participant count','235 distinct participants']], [2.1,4.8],8.5)
prompt(doc,'Confirm that this exact workbook is the file distributed by Dr. Prerna on Moodle. State the approved source in your final report.',2)

h(doc,'Dataset attributes',1)
attrs=csvrows('attribute_dictionary.csv')
table(doc,['Attribute','Meaning','Missing rows'],[[r['Attribute'],r['Description'],r['Missing_in_workbook']] for r in attrs],[1.1,4.6,1.2],7.6)
h(doc,'Data preprocessing',1)
audit=csvrows('data_audit.csv')
table(doc,['Audit measure','Count'],[[r['Measure'],r['Count']] for r in audit],[4.8,1.3],8)
fact(doc,'Target encoding','CDR = 0 is coded 0; CDR > 0 is coded 1; unknown CDR values remain excluded.')
fact(doc,'Predictor handling','ID and CDR are not predictors. Hand is constant. Delay is unavailable for target-known records. Nineteen eligible SES values are imputed using the training-set median.')
narrative(doc,'Rows without a recorded CDR cannot be assigned a reliable supervised target, so they are excluded rather than having labels inferred. This leaves 235 eligible rows: 135 with CDR equal to 0 and 100 with CDR above 0. Nineteen eligible SES values are missing; the median is estimated from training data and applied to the other partitions to prevent information from validation or test records entering preprocessing. ID is an identifier, not a clinical predictor. Hand has no variation, and Delay is unavailable for target-known records, so neither contributes to the models.')

h(doc,'Proposed methodology',1)
table(doc,['Stage','Project implementation'],[['Target','Binary recorded CDR group: 0 versus > 0'],['Split','Stratified 60/20/20 train, validation, and test; fixed seed 20261010'],['Feature blocks','Demographics; demographics + MRI (eTIV, nWBV); demographics + MRI + MMSE'],['Models','Logistic regression, shallow rpart tree, KNN (k=9), soft-voting average, majority baseline'],['Threshold','Choose on validation data by maximum F2; evaluate on test data'],['Metrics','ROC-AUC, accuracy, precision, recall, specificity, F1, F2, confusion matrix']],[1.5,5.4],8.3)
h(doc,'Novelty and improvement',1)
fact(doc,'Proposed contribution','Staged feature-block comparison plus validation-set threshold selection and soft voting, evaluated against demographic and majority-class baselines.')
narrative(doc,'The project contribution is a structured comparison rather than a claim of a novel clinical algorithm. It evaluates progressively richer feature blocks, compares several model families, selects thresholds using validation data, and tests a soft-voting candidate against demographic and majority-class baselines. This makes it possible to discuss whether added measures change test performance and to show the sensitivity-specificity trade-off. The method is an applied analytical framework for this dataset.')
h(doc,'Method limitation to discuss',2)
doc.add_paragraph('MMSE is a contemporaneous cognitive score and may overlap conceptually with CDR. Results from the MMSE-augmented model do not establish future cognitive decline prediction. Compare demographic-only and MRI-only results separately.')

h(doc,'Implementation and reproducibility',1)
table(doc,['Tool','Version or role'],[['R','4.3.2; preprocessing, models, evaluation, plots'],['rpart','4.1.21; shallow classification tree'],['Python','3.12.14; XLSX conversion'],['openpyxl','3.1.5; read-only workbook conversion']],[1.7,5.2],9)
fact(doc,'Run commands','python prepare_dataset.py; then Rscript analysis.R')
fact(doc,'Random seed','20261010')
fact(doc,'Source control','Add your GitHub repository URL after you create and review your repository.')
narrative(doc,'The Python preparation script converts the provided workbook sheet to CSV. The R script audits the data, excludes unknown targets, performs the stratified train-validation-test split, learns preprocessing from training data, fits the candidate models, selects thresholds using validation F2, and calculates held-out test metrics and plots. The fixed random seed is 20261010. Re-running the scripts from the project directory reproduces the outputs, subject to the stated software versions and input workbook.')
h(doc,'Performance evaluation',1)
metrics=csvrows('model_comparison_test.csv')
table(doc,['Model','Thr.','AUC','Acc.','Recall','Spec.','F2'],[[r['Model'],f"{float(r['Threshold']):.2f}",f"{float(r['AUC']):.3f}",f"{float(r['Accuracy']):.3f}",f"{float(r['Recall']):.3f}",f"{float(r['Specificity']):.3f}",f"{float(r['F2']):.3f}"] for r in metrics],[2.5,.5,.6,.6,.7,.7,.5],7.1)
doc.add_paragraph('Thresholds for candidate models were selected on validation data. The majority-class baseline uses 0.50. Test results use 47 records: 27 CDR 0 and 20 CDR > 0.')

h(doc,'Experimental results and visualizations',1)
image(doc,'01_target_availability.png',5.8,'Figure 1. Known target groups and rows excluded because CDR is missing.')
image(doc,'04_test_model_metrics.png',5.8,'Figure 2. Held-out test AUC and F2 for candidate models.')
narrative(doc,'On the held-out test set, the soft-voting ensemble has the highest ROC-AUC at 0.893, closely followed by the shallow tree at 0.892. The tree has the highest F2 at 0.893. These measures answer different questions: AUC summarizes ranking across thresholds, while F2 emphasizes recall at the selected threshold. KNN has the highest accuracy (0.766) and specificity (0.741), but not the highest recall. No model leads every metric, so selection depends on the intended error trade-off.')

h(doc,'Confusion matrix and healthcare insights',1)
image(doc,'05_ensemble_confusion_matrix.png',4.8,'Figure 3. Soft-voting ensemble at its validation-selected threshold.')
fact(doc,'Ensemble test result','At threshold 0.17: TN 16, FP 11, FN 2, TP 18; recall 0.900 and specificity 0.593.')
image(doc,'03_mmse_nwbv_by_cdr.png',5.7,'Figure 4. MMSE and nWBV distributions by recorded CDR group.')
narrative(doc,'Among records with a known CDR, the median age is 71 years in the CDR 0 group and 77 years in the CDR above 0 group. Median MMSE is 29 versus 26, and median nWBV is 0.773 versus 0.723. These summaries show group-level differences, but distributions overlap and do not establish causation. MMSE is contemporaneous with CDR, and the observed sample is limited to 235 labeled records from a workbook with 436 rows.')

h(doc,'Comparison with baseline and discussion',1)
fact(doc,'Demographic logistic baseline','Test AUC 0.719; test F2 0.763.')
fact(doc,'MRI-augmented logistic','Test AUC 0.754; test F2 0.812.')
fact(doc,'MMSE-augmented logistic','Test AUC 0.889; test F2 0.885. This is contemporaneous classification.')
fact(doc,'Top test metrics','Soft voting had AUC 0.893. The shallow tree had F2 0.893. KNN had accuracy 0.766 and specificity 0.741.')
h(doc,'Discussion',2); narrative(doc,'Adding MRI measures to the demographic logistic model increases test AUC from 0.719 to 0.754 and F2 from 0.763 to 0.812 in this split. The MMSE-augmented logistic model reaches AUC 0.889 and F2 0.885, but this gain needs caution because MMSE is a contemporaneous cognitive measure that may overlap with the construct represented by CDR. At its validation-selected threshold, the ensemble identifies 18 of 20 positive test records (recall 0.900) and correctly identifies 16 of 27 CDR 0 records (specificity 0.593), while producing 11 false positives and 2 false negatives. A lower threshold favors sensitivity but increases false alarms. The test set has only 47 records, so a few changed predictions could materially alter the metrics. In addition, 201 records have missing CDR; their exclusion may create selection bias if target availability is related to participant characteristics. These findings describe one split and do not establish future prediction or clinical usefulness.')
h(doc,'Conclusion',2); narrative(doc,'In this analysis, demographic, MRI, and especially contemporaneous MMSE features distinguish recorded CDR groups to varying degrees on the held-out split. The ensemble and shallow tree achieve similar AUC, while the tree leads on F2 and KNN leads on accuracy and specificity. The result supports comparing feature blocks and model trade-offs; it does not establish a clinically validated classifier or predict later dementia. The small test set, substantial missingness in CDR, and cross-sectional design limit generalization.')

h(doc,'Limitations',1)
for item in ['201 of 436 rows have missing CDR and are excluded; the analyzed subset may not represent the full workbook.','One stratified split gives uncertain estimates on a 47-record test set.','MMSE is contemporaneous with CDR and does not support future-risk claims.','Approval of the supplied workbook as the Moodle dataset must be confirmed.','The analysis uses summarized MRI measures, not raw MRI scans, and is not a clinical system.']: doc.add_paragraph(item,style='List Bullet')
h(doc,'Future scope',1); narrative(doc,'The next evaluation should use repeated stratified cross-validation or nested resampling to estimate variability and reduce dependence on one split. Calibration analysis would test whether predicted probabilities correspond to observed frequencies. Subgroup checks could identify uneven performance, provided group sizes are adequate and privacy is protected. The missing-CDR records should be examined to understand whether labeled and unlabeled rows differ. An independent approved dataset would be needed to assess external validity. Any extension should retain the distinction between cross-sectional classification and prediction of a future outcome.')
h(doc,'References',1)
doc.add_paragraph('Marcus, D. S., Wang, T. H., Parker, J., Csernansky, J. G., Morris, J. C., and Buckner, R. L. (2007). Open Access Series of Imaging Studies (OASIS): Cross-sectional MRI data in young, middle-aged, nondemented, and demented older adults. Journal of Cognitive Neuroscience, 19(9), 1498-1507. https://doi.org/10.1162/jocn.2007.19.9.1498')
doc.add_paragraph('OASIS Brains. Official OASIS project site. https://sites.wustl.edu/oasisbrains/')
prompt(doc,'Add any approved course materials and sources you personally consulted. Apply your instructor’s required citation style.',2)

doc.core_properties.title='Cognitive Status Classification Using OASIS-1 Clinical and MRI Measures'
doc.core_properties.subject='Healthcare Analytics Mini Project working report scaffold'
doc.core_properties.author='Atharv Jaiswal and Shreya Bansal'
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
