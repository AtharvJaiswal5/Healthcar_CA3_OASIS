"""Convert the supplied OASIS workbook sheet to a portable CSV."""
import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="data/oasis_cross-sectional-5708aa0a98d82080.xlsx")
    parser.add_argument("--output", default="data/oasis_cross-sectional.csv")
    args = parser.parse_args()
    source = Path(args.input)
    destination = Path(args.output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    with destination.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])
    workbook.close()
    print(f"Wrote {destination} from sheet {worksheet.title!r}.")


if __name__ == "__main__":
    main()
